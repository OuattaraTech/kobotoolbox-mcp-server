#!/usr/bin/env python3
"""
Renders an analysis report produced from KoboToolbox data into finished
deliverables: an analytical Excel workbook (native, editable charts), a Word
report, and/or a PDF.

Reads a JSON "report spec" from a file path given as argv[1] and prints a JSON
result to stdout describing the files that were written.

The spec is authored by the MCP server (structure) and by the model (narrative
text, which tables and charts matter for the stated objective).
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import xlsxwriter

# --------------------------------------------------------------------------
# Shared look and feel
# --------------------------------------------------------------------------

# Colour-blind-safe qualitative palette, readable on white in print.
PALETTE = [
    "#2F6FAF", "#E1812C", "#3A923A", "#C03D3E", "#9372B2",
    "#8B5B4C", "#D684BD", "#7F7F7F", "#BCBD45", "#4BA3C3",
]

MAX_PIE_SLICES = 8
MAX_CHART_CATEGORIES = 25
EXCEL_CELL_LIMIT = 32767


def sheet_name(raw: str, used: set[str]) -> str:
    """Excel sheet names: <=31 chars, no []:*?/\\, and unique in the workbook."""
    clean = re.sub(r"[\[\]:*?/\\]", "-", str(raw)).strip() or "Feuille"
    clean = clean[:31]
    candidate = clean
    i = 2
    while candidate.lower() in used:
        suffix = f" ({i})"
        candidate = clean[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate


def cell_value(v):
    """Flattens a JSON value into something a spreadsheet cell can hold."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Oui" if v else "Non"
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    s = str(v)
    return s[:EXCEL_CELL_LIMIT]


def setup_print(ws) -> None:
    """Makes a sheet print sensibly: landscape, scaled to one page wide."""
    ws.set_landscape()
    ws.set_paper(9)  # A4
    ws.fit_to_pages(1, 0)
    ws.set_margins(0.5, 0.5, 0.6, 0.5)


# --------------------------------------------------------------------------
# Excel deliverable
# --------------------------------------------------------------------------


def build_xlsx(spec: dict, path: str) -> None:
    wb = xlsxwriter.Workbook(path, {"nan_inf_to_errors": True, "constant_memory": False})

    fmt = {
        "title": wb.add_format({"bold": True, "font_size": 18, "font_color": "#1F3864"}),
        "subtitle": wb.add_format({"font_size": 11, "font_color": "#555555", "italic": True}),
        "h2": wb.add_format({"bold": True, "font_size": 13, "font_color": "#1F3864", "bottom": 1, "border_color": "#B4C6E7"}),
        "label": wb.add_format({"bold": True, "font_color": "#444444", "valign": "top"}),
        "body": wb.add_format({"text_wrap": True, "valign": "top", "align": "left"}),
        "bullet": wb.add_format({"text_wrap": True, "valign": "top", "align": "left", "indent": 1}),
        "th": wb.add_format({"bold": True, "bg_color": "#1F3864", "font_color": "white", "border": 1, "border_color": "#B4C6E7", "text_wrap": True, "valign": "vcenter"}),
        "td": wb.add_format({"border": 1, "border_color": "#D9D9D9", "valign": "top"}),
        "td_num": wb.add_format({"border": 1, "border_color": "#D9D9D9", "num_format": "#,##0.##"}),
        "td_tot": wb.add_format({"border": 1, "border_color": "#D9D9D9", "bold": True, "bg_color": "#EDF2F9", "num_format": "#,##0.##"}),
        "note": wb.add_format({"font_size": 9, "italic": True, "font_color": "#666666", "text_wrap": True, "valign": "top", "align": "left"}),
        "warn": wb.add_format({"font_size": 10, "font_color": "#9C0006", "bg_color": "#FFC7CE", "text_wrap": True, "valign": "top", "align": "left"}),
    }

    used_names: set[str] = set()

    # ---- Sheet 1: executive summary -------------------------------------
    ws = wb.add_worksheet(sheet_name("Synthèse", used_names))
    ws.hide_gridlines(2)
    setup_print(ws)
    ws.set_column("A:A", 3)
    ws.set_column("B:B", 26)
    ws.set_column("C:H", 16)

    r = 1
    ws.write(r, 1, spec.get("title") or "Rapport d'analyse", fmt["title"])
    r += 1
    generated = spec.get("generated_at") or datetime.now().isoformat(timespec="seconds")
    ws.write(r, 1, f"Généré le {generated}", fmt["subtitle"])
    r += 2

    meta_rows = [
        ("Objectif de l'analyse", spec.get("objective", "")),
        ("Formulaire source", spec.get("form_name", "")),
        ("Soumissions analysées", spec.get("n_rows", "")),
        ("Méthodologie", spec.get("methodology", "")),
    ]
    for label, value in meta_rows:
        if value in ("", None):
            continue
        ws.write(r, 1, label, fmt["label"])
        ws.merge_range(r, 2, r, 7, cell_value(value), fmt["body"])
        ws.set_row(r, max(15, min(90, 15 * (len(str(value)) // 90 + 1))))
        r += 1
    r += 1

    def write_block(heading: str, items: list, bullet: bool = True) -> None:
        nonlocal r
        if not items:
            return
        ws.merge_range(r, 1, r, 7, heading, fmt["h2"])
        r += 1
        for item in items:
            text = ("- " + str(item)) if bullet else str(item)
            ws.merge_range(r, 1, r, 7, text, fmt["bullet"] if bullet else fmt["body"])
            ws.set_row(r, max(15, min(120, 15 * (len(text) // 100 + 1))))
            r += 1
        r += 1

    if spec.get("summary"):
        write_block("Résumé exécutif", [spec["summary"]], bullet=False)
    write_block("Principaux constats", spec.get("findings") or [])
    write_block("Recommandations", spec.get("recommendations") or [])

    quality_notes = (spec.get("quality") or {}).get("notes") or []
    if quality_notes:
        ws.merge_range(r, 1, r, 7, "Limites et qualité des données", fmt["h2"])
        r += 1
        for note in quality_notes:
            ws.merge_range(r, 1, r, 7, "- " + str(note), fmt["warn"])
            ws.set_row(r, max(15, min(120, 15 * (len(str(note)) // 100 + 1))))
            r += 1

    # ---- Section sheets: tables + native charts --------------------------
    for section in spec.get("sections") or []:
        heading = section.get("heading") or "Analyse"
        sws = wb.add_worksheet(sheet_name(heading, used_names))
        sws.hide_gridlines(2)
        setup_print(sws)
        sws.set_column("A:A", 3)
        sws.set_column("B:B", 34)
        sws.set_column("C:N", 15)

        row = 1
        sws.merge_range(row, 1, row, 8, heading, fmt["title"])
        row += 2

        if section.get("text"):
            for para in str(section["text"]).split("\n"):
                if not para.strip():
                    row += 1
                    continue
                sws.merge_range(row, 1, row, 8, para.strip(), fmt["body"])
                sws.set_row(row, max(15, min(150, 15 * (len(para) // 110 + 1))))
                row += 1
            row += 1

        for item in section.get("items") or []:
            if item.get("block") == "chart":
                row = write_chart(wb, sws, item, row, fmt)
            else:
                row = write_table(sws, item, row, fmt)
                row += 1

    # ---- Clean data sheet, as a real Excel Table --------------------------
    dataset = spec.get("dataset") or {}
    columns = dataset.get("columns") or []
    rows = dataset.get("rows") or []
    if columns:
        dws = wb.add_worksheet(sheet_name("Données nettoyées", used_names))
        headers = [c.get("header") or c.get("path") or f"col{i}" for i, c in enumerate(columns)]

        # Excel Tables reject duplicate headers.
        seen: dict[str, int] = {}
        unique_headers = []
        for h in headers:
            h = str(h)[:255]
            if h.lower() in seen:
                seen[h.lower()] += 1
                h = f"{h} ({seen[h.lower()]})"
            else:
                seen[h.lower()] = 1
            unique_headers.append(h)

        hint = (
            "Ces données sont nettoyées et libellées. Pour créer un tableau croisé dynamique : "
            "cliquez dans le tableau, puis Insertion > Tableau croisé dynamique (la plage "
            "'DonneesKobo' se remplit automatiquement)."
        )
        dws.merge_range(0, 0, 0, max(len(unique_headers) - 1, 1), hint, fmt["note"])

        data_matrix = []
        for rec in rows:
            data_matrix.append([cell_value(rec.get(c.get("path"))) for c in columns])

        dws.add_table(
            2,
            0,
            2 + max(len(data_matrix), 1),
            len(unique_headers) - 1,
            {
                "name": "DonneesKobo",
                "style": "Table Style Medium 2",
                "columns": [{"header": h} for h in unique_headers],
                "data": data_matrix or [["" for _ in unique_headers]],
            },
        )
        for i, h in enumerate(unique_headers):
            dws.set_column(i, i, min(max(12, len(h) + 2), 40))
        dws.freeze_panes(3, 0)

    # ---- Data quality sheet ----------------------------------------------
    quality = spec.get("quality") or {}
    qcols = quality.get("columns") or []
    if qcols:
        qws = wb.add_worksheet(sheet_name("Qualité des données", used_names))
        qws.hide_gridlines(2)
        setup_print(qws)
        qws.set_column("A:A", 3)
        qws.set_column("B:B", 40)
        qws.set_column("C:G", 16)
        row = 1
        qws.merge_range(row, 1, row, 6, "Qualité des données", fmt["title"])
        row += 2
        summary_bits = [
            f"Soumissions analysées : {quality.get('rowsFetched', '?')} sur {quality.get('rowsTotal', '?')}",
            f"Doublons potentiels : {quality.get('duplicateRows', 0)}",
        ]
        for bit in summary_bits:
            qws.merge_range(row, 1, row, 6, bit, fmt["body"])
            row += 1
        row += 1
        table = {
            "columns": ["Question", "Type", "Réponses manquantes", "% manquant", "Valeurs distinctes"],
            "rows": [
                [
                    c.get("header"),
                    c.get("measure"),
                    c.get("missing"),
                    c.get("missingPct"),
                    c.get("unique"),
                ]
                for c in qcols
            ],
        }
        write_table(qws, table, row, fmt)

    wb.close()


def write_table(ws, table: dict, row: int, fmt: dict) -> int:
    """Writes a titled, bordered table and returns the next free row."""
    title = table.get("title")
    if title:
        ws.merge_range(row, 1, row, 8, title, fmt["h2"])
        row += 1

    cols = table.get("columns") or []
    if not cols:
        return row

    for ci, header in enumerate(cols):
        ws.write(row, 1 + ci, cell_value(header), fmt["th"])
    ws.set_row(row, 30)
    row += 1

    total_label = str(table.get("total_row_label") or "").strip().lower()
    for rec in table.get("rows") or []:
        is_total = bool(rec) and str(rec[0]).strip().lower() == total_label and total_label
        for ci, value in enumerate(rec):
            v = cell_value(value)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                ws.write_number(row, 1 + ci, v, fmt["td_tot"] if is_total else fmt["td_num"])
            else:
                ws.write(row, 1 + ci, v, fmt["td_tot"] if is_total else fmt["td"])
        row += 1

    if table.get("note"):
        ws.merge_range(row, 1, row, 8, str(table["note"]), fmt["note"])
        ws.set_row(row, max(15, min(90, 15 * (len(str(table["note"])) // 110 + 1))))
        row += 1
    return row


def write_chart(wb, ws, chart: dict, row: int, fmt: dict) -> int:
    """
    Writes the chart's source data into the sheet, then inserts a *native*
    Excel chart bound to those cells - so it stays editable and refreshable.
    """
    categories = [str(c) for c in (chart.get("categories") or [])][:MAX_CHART_CATEGORIES]
    series = chart.get("series") or []
    if not categories or not series:
        return row

    kind = (chart.get("kind") or "column").lower()
    if kind not in ("column", "bar", "pie", "line", "doughnut", "area", "scatter"):
        kind = "column"
    if kind == "pie" and len(categories) > MAX_PIE_SLICES:
        kind = "bar"

    title = chart.get("title") or ""
    if title:
        ws.merge_range(row, 1, row, 8, title, fmt["h2"])
        row += 1

    # Source data block. It is left visible on purpose: the reader can see the
    # numbers behind the picture, and Excel keeps the chart linked to them.
    data_start = row
    ws.write(row, 1, "Catégorie", fmt["th"])
    for si, s in enumerate(series):
        ws.write(row, 2 + si, cell_value(s.get("name") or f"Série {si + 1}"), fmt["th"])
    row += 1

    for ci, cat in enumerate(categories):
        ws.write(row, 1, cat, fmt["td"])
        for si, s in enumerate(series):
            values = s.get("values") or []
            v = values[ci] if ci < len(values) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                ws.write_number(row, 2 + si, v, fmt["td_num"])
            else:
                ws.write_blank(row, 2 + si, None, fmt["td_num"])
        row += 1
    data_end = row - 1

    excel_chart = wb.add_chart({"type": kind})
    name = ws.get_name()
    for si, s in enumerate(series):
        color = PALETTE[si % len(PALETTE)]
        opts = {
            "name": [name, data_start, 2 + si],
            "categories": [name, data_start + 1, 1, data_end, 1],
            "values": [name, data_start + 1, 2 + si, data_end, 2 + si],
        }
        # Keys are added only when they apply: xlsxwriter rejects several of
        # these when set to a falsy placeholder.
        if kind == "line":
            opts["line"] = {"color": color, "width": 2.25}
            opts["marker"] = {"type": "circle", "size": 5, "fill": {"color": color}}
        else:
            opts["fill"] = {"color": color}
            opts["border"] = {"none": True}
        if kind in ("pie", "doughnut"):
            opts["data_labels"] = {"percentage": True, "font": {"size": 8}}
            opts["points"] = [{"fill": {"color": PALETTE[i % len(PALETTE)]}} for i in range(len(categories))]
        excel_chart.add_series(opts)

    excel_chart.set_title({"name": title or " ", "overlay": False})
    if kind not in ("pie", "doughnut"):
        excel_chart.set_x_axis({"name": chart.get("x_title") or "", "num_font": {"size": 9}})
        excel_chart.set_y_axis({"name": chart.get("y_title") or "", "major_gridlines": {"visible": True}})
    excel_chart.set_legend({"position": "bottom" if len(series) > 1 or kind in ("pie", "doughnut") else "none"})
    excel_chart.set_style(2)
    excel_chart.set_size({"width": 620, "height": 360})

    ws.insert_chart(data_start, 2 + len(series) + 1, excel_chart)

    # Leave room so the floating chart doesn't overlap what comes next.
    row = max(row, data_start + 19)
    if chart.get("note"):
        ws.merge_range(row, 1, row, 8, str(chart["note"]), fmt["note"])
        row += 1
    return row + 1


# --------------------------------------------------------------------------
# Chart images (for Word / PDF)
# --------------------------------------------------------------------------


def render_chart_png(chart: dict, out_path: str) -> str | None:
    categories = [str(c) for c in (chart.get("categories") or [])][:MAX_CHART_CATEGORIES]
    series = chart.get("series") or []
    if not categories or not series:
        return None

    kind = (chart.get("kind") or "column").lower()
    if kind == "pie" and len(categories) > MAX_PIE_SLICES:
        kind = "bar"

    def nums(s):
        vals = s.get("values") or []
        out = []
        for i in range(len(categories)):
            v = vals[i] if i < len(vals) else None
            out.append(float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0)
        return out

    fig_h = max(3.2, min(8.0, 0.32 * len(categories) + 1.8)) if kind == "bar" else 4.0
    fig, ax = plt.subplots(figsize=(7.2, fig_h), dpi=160)

    if kind in ("pie", "doughnut"):
        values = nums(series[0])
        total = sum(values)
        wedges, _texts, autotexts = ax.pie(
            values,
            labels=categories,
            colors=[PALETTE[i % len(PALETTE)] for i in range(len(categories))],
            autopct=(lambda p: f"{p:.1f}%") if total else None,
            startangle=90,
            counterclock=False,
            textprops={"fontsize": 9},
            wedgeprops={"width": 0.55, "edgecolor": "white"} if kind == "doughnut" else {"edgecolor": "white"},
        )
        for t in autotexts:
            t.set_fontsize(8)
            t.set_color("white")
        ax.axis("equal")
    elif kind == "line":
        for si, s in enumerate(series):
            ax.plot(categories, nums(s), marker="o", linewidth=2,
                    color=PALETTE[si % len(PALETTE)], label=str(s.get("name") or f"Serie {si+1}"))
        ax.grid(axis="y", color="#E6E6E6")
        ax.set_axisbelow(True)
        plt.setp(ax.get_xticklabels(), rotation=30, ha="right", fontsize=8)
    elif kind == "bar":
        n = len(series)
        height = 0.8 / n
        ypos = range(len(categories))
        for si, s in enumerate(series):
            offs = [y + si * height - (0.8 - height) / 2 for y in ypos]
            ax.barh(offs, nums(s), height=height,
                    color=PALETTE[si % len(PALETTE)], label=str(s.get("name") or f"Serie {si+1}"))
        ax.set_yticks(list(ypos))
        ax.set_yticklabels(categories, fontsize=8)
        ax.invert_yaxis()
        ax.grid(axis="x", color="#E6E6E6")
        ax.set_axisbelow(True)
    else:  # column
        n = len(series)
        width = 0.8 / n
        xpos = range(len(categories))
        for si, s in enumerate(series):
            offs = [x + si * width - (0.8 - width) / 2 for x in xpos]
            ax.bar(offs, nums(s), width=width,
                   color=PALETTE[si % len(PALETTE)], label=str(s.get("name") or f"Serie {si+1}"))
        ax.set_xticks(list(xpos))
        ax.set_xticklabels(categories, rotation=30, ha="right", fontsize=8)
        ax.grid(axis="y", color="#E6E6E6")
        ax.set_axisbelow(True)

    if chart.get("title"):
        ax.set_title(str(chart["title"]), fontsize=11, fontweight="bold", color="#1F3864")
    if kind not in ("pie", "doughnut"):
        if chart.get("x_title"):
            ax.set_xlabel(str(chart["x_title"]), fontsize=9)
        if chart.get("y_title"):
            ax.set_ylabel(str(chart["y_title"]), fontsize=9)
        for spine in ("top", "right"):
            ax.spines[spine].set_visible(False)
        if len(series) > 1:
            ax.legend(fontsize=8, frameon=False)

    fig.tight_layout()
    fig.savefig(out_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_path


# --------------------------------------------------------------------------
# Word deliverable
# --------------------------------------------------------------------------


def build_docx(spec: dict, path: str, tmpdir: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def heading(text: str, level: int) -> None:
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    def para(text: str, italic: bool = False, size: int = 11, color=None) -> None:
        p = doc.add_paragraph()
        run = p.add_run(str(text))
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color

    # ---- Title page ----
    title = doc.add_heading(spec.get("title") or "Rapport d'analyse", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    srun = sub.add_run(
        f"{spec.get('form_name', '')}\nGénéré le {spec.get('generated_at') or datetime.now().isoformat(timespec='seconds')}"
    )
    srun.italic = True
    srun.font.size = Pt(10)
    srun.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if spec.get("objective"):
        heading("Objectif de l'analyse", 1)
        para(spec["objective"])

    if spec.get("methodology"):
        heading("Méthodologie", 1)
        para(spec["methodology"])
    if spec.get("n_rows"):
        para(f"Base d'analyse : {spec['n_rows']} soumission(s).", italic=True, size=10)

    if spec.get("summary"):
        heading("Résumé exécutif", 1)
        para(spec["summary"])

    if spec.get("findings"):
        heading("Principaux constats", 1)
        for f in spec["findings"]:
            doc.add_paragraph(str(f), style="List Bullet")

    # ---- Analysis sections ----
    chart_idx = 0
    for section in spec.get("sections") or []:
        heading(section.get("heading") or "Analyse", 1)
        if section.get("text"):
            for block in str(section["text"]).split("\n"):
                if block.strip():
                    para(block.strip())

        for item in section.get("items") or []:
            if item.get("block") != "chart":
                add_docx_table(doc, item, Pt)
                continue
            png = os.path.join(tmpdir, f"chart_{chart_idx}.png")
            chart_idx += 1
            if render_chart_png(item, png):
                doc.add_picture(png, width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if item.get("note"):
                    cap = doc.add_paragraph()
                    crun = cap.add_run(str(item["note"]))
                    crun.italic = True
                    crun.font.size = Pt(9)
                    crun.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if spec.get("recommendations"):
        heading("Recommandations", 1)
        for rec in spec["recommendations"]:
            doc.add_paragraph(str(rec), style="List Bullet")

    # ---- Quality annex ----
    quality = spec.get("quality") or {}
    if quality:
        doc.add_page_break()
        heading("Annexe — Qualité des données", 1)
        para(
            f"Soumissions analysées : {quality.get('rowsFetched', '?')} sur {quality.get('rowsTotal', '?')}. "
            f"Doublons potentiels : {quality.get('duplicateRows', 0)}.",
        )
        for note in quality.get("notes") or []:
            doc.add_paragraph(str(note), style="List Bullet")
        qcols = quality.get("columns") or []
        if qcols:
            add_docx_table(
                doc,
                {
                    "title": "Complétude par question",
                    "columns": ["Question", "Manquantes", "% manquant", "Valeurs distinctes"],
                    "rows": [
                        [c.get("header"), c.get("missing"), c.get("missingPct"), c.get("unique")]
                        for c in qcols
                    ],
                },
                Pt,
            )

    doc.save(path)


def add_docx_table(doc, table: dict, Pt):
    from docx.shared import RGBColor

    cols = table.get("columns") or []
    if not cols:
        return 0
    if table.get("title"):
        p = doc.add_paragraph()
        run = p.add_run(str(table["title"]))
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    rows = table.get("rows") or []
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, c in enumerate(cols):
        hdr[i].text = str(c)
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    for rec in rows:
        cells = t.add_row().cells
        for i in range(len(cols)):
            v = rec[i] if i < len(rec) else ""
            cells[i].text = "" if v is None else str(cell_value(v))
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    if table.get("note"):
        p = doc.add_paragraph()
        run = p.add_run(str(table["note"]))
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return 0


# --------------------------------------------------------------------------
# PDF (via LibreOffice)
# --------------------------------------------------------------------------

# Only the Linux package managers put LibreOffice on the PATH. The macOS and
# Windows installers drop it at a fixed location and leave the PATH alone, so
# those have to be probed explicitly before giving up.
SOFFICE_INSTALL_PATHS = {
    "darwin": [
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/opt/homebrew/bin/soffice",
        "/usr/local/bin/soffice",
    ],
    "win32": [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ],
    "linux": [
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        "/snap/bin/libreoffice",
        "/var/lib/flatpak/exports/bin/org.libreoffice.LibreOffice",
    ],
}


def soffice_install_hint() -> str:
    """The install command to quote back when LibreOffice is missing."""
    if sys.platform == "darwin":
        return "brew install --cask libreoffice"
    if sys.platform == "win32":
        return "installeur à télécharger sur https://www.libreoffice.org/download/"
    return "sudo apt install libreoffice"


def find_soffice() -> str | None:
    """Locates the LibreOffice binary used for PDF conversion, or None."""
    override = os.environ.get("SOFFICE_BIN", "").strip()
    if override:
        if os.path.isfile(override):
            return override
        return shutil.which(override)

    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found

    candidates = list(SOFFICE_INSTALL_PATHS.get(sys.platform, []))
    if sys.platform == "win32":
        # Per-user installs land under %LOCALAPPDATA% rather than Program Files.
        local_app_data = os.environ.get("LOCALAPPDATA")
        if local_app_data:
            candidates.append(
                os.path.join(local_app_data, "Programs", "LibreOffice", "program", "soffice.exe")
            )
    for candidate in candidates:
        if os.path.isfile(candidate):
            return candidate
    return None


def build_pdf(docx_path: str, out_dir: str) -> str:
    """Converts the Word report to PDF. Requires LibreOffice on the machine."""
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice est introuvable : la conversion PDF est impossible "
            "(les sorties xlsx et docx, elles, ne sont pas concernées). "
            f"Installez-le ({soffice_install_hint()}) ou renseignez SOFFICE_BIN "
            "dans le .env avec le chemin complet du binaire."
        )
    profile = Path(tempfile.gettempdir()) / "kobo_mcp_lo_profile"
    cmd = [
        soffice,
        "--headless",
        "--norestore",
        # as_uri() rather than "file://" + path: on Windows the temp directory
        # carries a drive letter and backslashes, which are not a valid URL.
        f"-env:UserInstallation={profile.as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        out_dir,
        docx_path,
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    produced = os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    if not os.path.exists(produced):
        raise RuntimeError(
            f"La conversion PDF via LibreOffice ({soffice}) a échoué. "
            f"stdout={proc.stdout.strip()[:400]} stderr={proc.stderr.strip()[:400]}"
        )
    return produced


# --------------------------------------------------------------------------


def main() -> int:
    if len(sys.argv) >= 2 and sys.argv[1] == "--check-soffice":
        soffice = find_soffice()
        print(
            json.dumps(
                {"ok": bool(soffice), "path": soffice or "", "hint": soffice_install_hint()},
                ensure_ascii=False,
            )
        )
        return 0

    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "Usage: render_report.py <spec.json>"}))
        return 1

    with open(sys.argv[1], "r", encoding="utf-8") as fh:
        spec = json.load(fh)

    out_dir = spec.get("output_dir") or os.getcwd()
    os.makedirs(out_dir, exist_ok=True)
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", str(spec.get("base_name") or "rapport"))[:80]
    formats = [f.lower() for f in (spec.get("formats") or ["xlsx"])]

    files = []
    warnings = []

    if "xlsx" in formats:
        p = os.path.join(out_dir, base + ".xlsx")
        build_xlsx(spec, p)
        files.append({"format": "xlsx", "path": p, "size_bytes": os.path.getsize(p)})

    needs_docx = "docx" in formats or "pdf" in formats
    if needs_docx:
        with tempfile.TemporaryDirectory() as tmpdir:
            docx_path = os.path.join(out_dir if "docx" in formats else tmpdir, base + ".docx")
            build_docx(spec, docx_path, tmpdir)
            if "docx" in formats:
                files.append({"format": "docx", "path": docx_path, "size_bytes": os.path.getsize(docx_path)})
            if "pdf" in formats:
                try:
                    pdf_path = build_pdf(docx_path, out_dir)
                    files.append({"format": "pdf", "path": pdf_path, "size_bytes": os.path.getsize(pdf_path)})
                except Exception as exc:  # noqa: BLE001 - reported back to the caller
                    warnings.append(str(exc))

    print(json.dumps({"ok": True, "files": files, "warnings": warnings}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001 - surfaced as a tool error
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
